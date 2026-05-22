from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path("Document") / "test-cases"
ISSUE_DATE = "21/05/2026"
PROJECT_NAME = "Quản Lý Nhà Hàng / Restaurant Management System"
PROJECT_CODE = "RMS-FLUTTER-NODE"
VERSION = "1.0"
PENDING = "Pending / Chưa chạy"
PASS = "Pass / Đạt"

all_registry: list[tuple[str, str, str]] = []
workbook_summaries: list[dict[str, object]] = []

header_fill = PatternFill("solid", fgColor="111827")
subheader_fill = PatternFill("solid", fgColor="D9EAF7")
meta_fill = PatternFill("solid", fgColor="EAF4FF")
pass_fill = PatternFill("solid", fgColor="DCFCE7")
fail_fill = PatternFill("solid", fgColor="FEE2E2")
pending_fill = PatternFill("solid", fgColor="FEF3C7")
white_font = Font(color="FFFFFF", bold=True)
bold_font = Font(bold=True)
thin = Side(style="thin", color="CBD5E1")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def c(
    id_: str,
    function: str,
    location: str,
    desc_vi: str,
    desc_en: str,
    expected_vi: str,
    expected_en: str,
    data: str = "",
    preconditions: str = "",
    result: str = PENDING,
    actual: str = "Chưa chạy / Not executed",
    note: str = "",
    test_date: str = "",
) -> dict[str, str]:
    step_vi = (
        f"1. Mở vị trí test: {location}.\n"
        f"2. Thực hiện tình huống: {desc_vi}.\n"
        "3. Kiểm tra giao diện, dữ liệu lưu trong hệ thống và phản hồi lỗi/thành công."
    )
    step_en = (
        f"1. Open test location: {location}.\n"
        f"2. Execute scenario: {desc_en}.\n"
        "3. Verify UI, saved data, and success/error response."
    )
    return {
        "id": id_,
        "function": function,
        "location": location,
        "desc_vi": desc_vi,
        "desc_en": desc_en,
        "preconditions": preconditions,
        "steps_vi": step_vi,
        "steps_en": step_en,
        "data": data,
        "expected_vi": expected_vi,
        "expected_en": expected_en,
        "test_date": test_date,
        "result": result,
        "actual": actual,
        "note": note,
    }


def build_cases(
    prefix: str,
    function: str,
    location: str,
    default_preconditions: str,
    items: list[tuple[str, ...]],
) -> list[dict[str, str]]:
    cases = []
    for index, item in enumerate(items, start=1):
        desc_vi, desc_en, expected_vi, expected_en, *rest = item
        data = rest[0] if len(rest) > 0 else ""
        pre = rest[1] if len(rest) > 1 and rest[1] else default_preconditions
        cases.append(
            c(
                f"{prefix}-{index:03d}",
                function,
                location,
                desc_vi,
                desc_en,
                expected_vi,
                expected_en,
                data=data,
                preconditions=pre,
            )
        )
    return cases


def add_cover(
    wb: Workbook,
    module_code: str,
    module_vi: str,
    module_en: str,
    requirement: str,
    total_cases: int,
) -> None:
    ws = wb.active
    ws.title = "Cover"
    ws.merge_cells("A1:H2")
    ws["A1"] = "TEST CASE"
    ws["A1"].font = Font(size=22, bold=True, color="FFFFFF")
    ws["A1"].fill = header_fill
    ws["A1"].alignment = center
    rows = [
        ("Project Name", PROJECT_NAME),
        ("Project Code", PROJECT_CODE),
        ("Module", f"{module_vi} / {module_en}"),
        ("Module Code", module_code),
        ("Version", VERSION),
        ("Issue date", ISSUE_DATE),
        ("Total test cases", total_cases),
        ("Language", "Tiếng Việt + English"),
        (
            "Result policy",
            "Pass/Fail only for executed smoke checks; Pending means written but not executed manually.",
        ),
    ]
    row = 4
    for label, value in rows:
        ws[f"A{row}"] = label
        ws[f"B{row}"] = value
        ws[f"A{row}"].font = bold_font
        ws[f"A{row}"].fill = meta_fill
        ws[f"A{row}"].border = border
        ws[f"B{row}"].border = border
        ws[f"B{row}"].alignment = wrap
        row += 1
    row += 1
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
    ws.cell(row, 1).value = "Record of change / Lịch sử thay đổi"
    ws.cell(row, 1).font = white_font
    ws.cell(row, 1).fill = header_fill
    ws.cell(row, 1).alignment = center
    row += 1
    headers = ["Version", "Date", "Author", "Change description", "Reviewer", "Status", "Note", ""]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row, col, h)
        cell.font = bold_font
        cell.fill = subheader_fill
        cell.border = border
        cell.alignment = center
    row += 1
    values = [VERSION, ISSUE_DATE, "Codex", f"Create bilingual test cases for {module_vi}", "", "Draft", requirement, ""]
    for col, v in enumerate(values, 1):
        cell = ws.cell(row, col, v)
        cell.border = border
        cell.alignment = wrap
    widths = [18, 32, 18, 44, 18, 16, 38, 14]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def add_test_cases_sheet(
    wb: Workbook,
    module_code: str,
    requirement: str,
    cases: list[dict[str, str]],
) -> None:
    ws = wb.create_sheet("Test Cases")
    max_formula_row = max(500, 12 + len(cases) + 20)
    ws.merge_cells("A1:O1")
    ws["A1"] = "TEST CASE"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = header_fill
    ws["A1"].alignment = center
    meta = [
        ("System Name:", PROJECT_NAME),
        ("Module Code:", module_code),
        ("Test requirement:", requirement),
        ("Pass", f'=COUNTIF(M12:M{max_formula_row},"Pass*")', "Pending", f'=COUNTIF(M12:M{max_formula_row},"Pending*")'),
        ("Fail", f'=COUNTIF(M12:M{max_formula_row},"Fail*")', "Number of test cases:", f"=COUNTA(A12:A{max_formula_row})"),
    ]
    for r, values in enumerate(meta, start=3):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(r, col, value)
            cell.border = border
            cell.alignment = wrap
            if col in (1, 3):
                cell.font = bold_font
                cell.fill = meta_fill
    headers = [
        "ID",
        "Function / Chức năng",
        "Test location / Vị trí test",
        "Test Case Description (VI)",
        "Test Case Description (EN)",
        "Preconditions / Tiền điều kiện",
        "Test Case Procedure (VI)",
        "Test Case Procedure (EN)",
        "Test Data / Dữ liệu",
        "Expected Output (VI)",
        "Expected Output (EN)",
        "Test date",
        "Result / Kết quả",
        "Actual Result / Kết quả thực tế",
        "Note / Ghi chú",
    ]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(9, col, h)
        cell.fill = header_fill
        cell.font = white_font
        cell.alignment = center
        cell.border = border
    start = 12
    keys = [
        "id",
        "function",
        "location",
        "desc_vi",
        "desc_en",
        "preconditions",
        "steps_vi",
        "steps_en",
        "data",
        "expected_vi",
        "expected_en",
        "test_date",
        "result",
        "actual",
        "note",
    ]
    for offset, case in enumerate(cases):
        row = start + offset
        for col, key in enumerate(keys, start=1):
            cell = ws.cell(row, col, case[key])
            cell.border = border
            cell.alignment = wrap if col not in (1, 12, 13) else center
        result_cell = ws.cell(row, 13)
        if str(case["result"]).startswith("Pass"):
            result_cell.fill = pass_fill
        elif str(case["result"]).startswith("Fail"):
            result_cell.fill = fail_fill
        else:
            result_cell.fill = pending_fill
    widths = [16, 24, 28, 34, 34, 34, 44, 44, 28, 38, 38, 14, 18, 38, 28]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A12"
    ws.auto_filter.ref = f"A9:O{start + len(cases) - 1}"
    for r in range(12, start + len(cases)):
        ws.row_dimensions[r].height = 96


def add_report_sheet(
    wb: Workbook,
    module_code: str,
    module_vi: str,
    module_en: str,
    cases: list[dict[str, str]],
) -> None:
    ws = wb.create_sheet("Test Report")
    ws.merge_cells("A1:G1")
    ws["A1"] = "TEST REPORT"
    ws["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    ws["A1"].fill = header_fill
    ws["A1"].alignment = center
    info = [
        ("System Name", PROJECT_NAME),
        ("Module", f"{module_vi} / {module_en}"),
        ("Module Code", module_code),
        ("Issue date", ISSUE_DATE),
        ("Total test cases", len(cases)),
    ]
    for row, (label, value) in enumerate(info, start=3):
        ws.cell(row, 1, label).font = bold_font
        ws.cell(row, 1).fill = meta_fill
        ws.cell(row, 2, value)
        for col in range(1, 3):
            ws.cell(row, col).border = border
            ws.cell(row, col).alignment = wrap
    headers = ["Result", "Count", "Meaning / Ý nghĩa", "Scope", "Owner", "Date", "Note"]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(10, col, h)
        cell.fill = header_fill
        cell.font = white_font
        cell.border = border
        cell.alignment = center
    rows = [
        ("Pass / Đạt", '=COUNTIF(\'Test Cases\'!M12:M500,"Pass*")', "Executed and matched expected result / Đã chạy và đúng mong đợi"),
        ("Fail / Không đạt", '=COUNTIF(\'Test Cases\'!M12:M500,"Fail*")', "Executed but did not match expected result / Đã chạy nhưng sai mong đợi"),
        ("Pending / Chưa chạy", '=COUNTIF(\'Test Cases\'!M12:M500,"Pending*")', "Written but not yet executed manually / Đã viết, chưa chạy thủ công"),
        ("Number of test cases", "=COUNTA('Test Cases'!A12:A500)", "Total case rows / Tổng số dòng case"),
    ]
    for idx, (result, formula, meaning) in enumerate(rows, start=11):
        values = [result, formula, meaning, module_vi, "QA / Dev", ISSUE_DATE, ""]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(idx, col, value)
            cell.border = border
            cell.alignment = wrap
        if result.startswith("Pass"):
            ws.cell(idx, 1).fill = pass_fill
        elif result.startswith("Fail"):
            ws.cell(idx, 1).fill = fail_fill
        elif result.startswith("Pending"):
            ws.cell(idx, 1).fill = pending_fill
    widths = [24, 14, 52, 28, 16, 14, 28]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def save_workbook(
    rel_path: str,
    module_code: str,
    module_vi: str,
    module_en: str,
    requirement: str,
    cases: list[dict[str, str]],
) -> None:
    path = ROOT / rel_path
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    add_cover(wb, module_code, module_vi, module_en, requirement, len(cases))
    add_test_cases_sheet(wb, module_code, requirement, cases)
    add_report_sheet(wb, module_code, module_vi, module_en, cases)
    wb.save(path)
    counts = Counter(case["result"].split(" / ")[0] for case in cases)
    workbook_summaries.append(
        {
            "path": str(path).replace("\\", "/"),
            "module_code": module_code,
            "module_vi": module_vi,
            "module_en": module_en,
            "count": len(cases),
            "pass": counts.get("Pass", 0),
            "fail": counts.get("Fail", 0),
            "pending": counts.get("Pending", 0),
        }
    )
    all_registry.extend((case["id"], str(path), case["desc_vi"]) for case in cases)


def executed_smoke_cases() -> list[dict[str, str]]:
    pre = "Các lệnh smoke đã chạy trong workspace ngày 21/05/2026 sau khi đổi password sang duyencocong2."
    function = "Smoke verification"
    location = "Local workspace / Backend 3000 / Web 8080 / Flutter"
    raw = [
        ("Kiểm tra cú pháp toàn bộ backend JS bằng node --check", "Check all backend JavaScript syntax with node --check", "Mọi file JS ngoài node_modules parse thành công.", "All JS files outside node_modules parse successfully.", "node --check", "node --check completed with exit code 0."),
        ("Gọi API root GET / ở cổng 3000", "Call API root GET / on port 3000", "API trả success true và message server đang chạy.", "API returns success true and running message.", "http://127.0.0.1:3000/", '{"success":true,"message":"Restaurant API is running"}'),
        ("Gọi API health GET /admin/health", "Call health API GET /admin/health", "API trả status ok và counts dữ liệu.", "API returns ok status and data counts.", "http://127.0.0.1:3000/admin/health", "status=ok, storage=json-file, tables=6, menuItems=5, employees=3."),
        ("Đăng nhập admin bằng password duyencocong2 qua API", "Login admin with password duyencocong2 through API", "API trả success true, role admin, không trả password.", "API returns success true, admin role, and no password.", "admin / duyencocong2", "success=true, user.role=admin, password omitted."),
        ("Chạy flutter analyze sau khi sửa password", "Run flutter analyze after password change", "Không có lỗi phân tích Dart/Flutter.", "No Dart/Flutter analysis issues.", "flutter analyze", "No issues found."),
        ("Chạy flutter test", "Run flutter test", "Widget test hiện có pass.", "Existing widget test passes.", "flutter test", "All tests passed."),
        ("Build web release", "Build release web", "Tạo build/web thành công.", "build/web is created successfully.", "flutter build web --release", "Built build\\web successfully."),
        ("Build Android APK debug", "Build Android debug APK", "Tạo app-debug.apk thành công.", "app-debug.apk is created successfully.", "flutter build apk --debug", "Built build\\app\\outputs\\flutter-apk\\app-debug.apk successfully."),
        ("Kiểm tra web server local cổng 8080", "Check local web server on port 8080", "Web server trả lời request HEAD.", "Web server responds to HEAD request.", "http://127.0.0.1:8080/", "web-server: ok"),
    ]
    cases = []
    for index, (vi, en, exp_vi, exp_en, data, actual) in enumerate(raw, start=1):
        cases.append(
            c(
                f"SMK-{index:03d}",
                function,
                location,
                vi,
                en,
                exp_vi,
                exp_en,
                data=data,
                preconditions=pre,
                result=PASS,
                actual=actual,
                test_date=ISSUE_DATE,
            )
        )
    return cases


def pending_specs() -> list[tuple[str, str, str, str, str, list[dict[str, str]]]]:
    specs = []
    mobile_pre = "Backend đang chạy cổng 3000; cài APK hoặc chạy app trên emulator/điện thoại; đăng nhập đúng vai trò."
    web_pre = "Backend đang chạy cổng 3000; web build chạy tại http://127.0.0.1:8080 hoặc flutter run -d chrome; đăng nhập đúng vai trò."
    api_pre = "Backend Node/Express đang chạy tại http://127.0.0.1:3000; dữ liệu JSON ở backend/data/restaurant.json."
    e2e_pre = "Backend, mobile app và web app đều dùng cùng dữ liệu; reset demo nếu cần trạng thái sạch."

    def add(rel, code, vi, en, req, function, location, pre, items):
        specs.append((rel, code, vi, en, req, build_cases(code, function, location, pre, items)))

    add("01_mobile_app/01_authentication/mobile_authentication_test_cases.xlsx", "MOB-AUTH", "Mobile - Đăng nhập", "Mobile - Authentication", "Kiểm tra đăng nhập, phân quyền và lỗi kết nối trên mobile", "Đăng nhập", "Mobile app > Login screen", mobile_pre, [
        ("Mở app lần đầu hiển thị logo và tên Quản Lý Nhà Hàng", "Launch app and display logo plus Restaurant Management title", "Màn hình đăng nhập hiển thị đúng thương hiệu, ô username, password và nút đăng nhập.", "Login screen shows branding, username/password fields, and sign-in button.", "Android APK debug"),
        ("Đăng nhập admin bằng mật khẩu duyencocong2", "Sign in as admin with password duyencocong2", "Đăng nhập thành công, chuyển vào dashboard và session có role admin.", "Login succeeds, dashboard opens, and session role is admin.", "admin / duyencocong2"),
        ("Đăng nhập staff bằng mật khẩu duyencocong2", "Sign in as staff with password duyencocong2", "Đăng nhập thành công, chuyển vào dashboard và session có role staff.", "Login succeeds, dashboard opens, and session role is staff.", "staff / duyencocong2"),
        ("Nhập sai mật khẩu admin", "Enter wrong admin password", "App hiển thị thông báo lỗi, không tạo session và vẫn ở màn hình đăng nhập.", "App shows error, does not create a session, and remains on login screen.", "admin / wrong-password"),
        ("Để trống username hoặc password", "Leave username or password empty", "App không crash, hiển thị lỗi hợp lệ từ backend hoặc thông báo nhập liệu.", "App does not crash and shows backend or validation error.", "empty fields"),
        ("Kiểm tra password được che khi nhập", "Verify password field is obscured while typing", "Ký tự mật khẩu không hiển thị trực tiếp trên màn hình.", "Password characters are not displayed in plain text.", "duyencocong2"),
        ("Đăng xuất khỏi tài khoản admin", "Log out from admin account", "Session bị xóa và app quay về màn hình đăng nhập.", "Session is cleared and app returns to login screen.", "admin session"),
        ("Tài khoản admin nhìn thấy nhóm chức năng quản trị", "Admin account can see management modules", "Admin thấy quản lý nhân viên, thông báo, tính lương và web admin.", "Admin sees employees, notifications, payroll, and web admin modules.", "role=admin"),
        ("Tài khoản staff bị ẩn chức năng chỉ dành cho admin", "Staff account hides admin-only modules", "Staff không thấy mục quản lý nhân viên, tính lương hoặc web admin.", "Staff does not see employees, payroll, or web admin modules.", "role=staff"),
        ("Backend tắt hoặc sai địa chỉ API", "Backend is stopped or API base URL is wrong", "App hiển thị lỗi kết nối dễ hiểu và không bị treo loading vô hạn.", "App shows a clear connection error and does not hang forever.", "API_BASE_URL invalid"),
        ("Bàn phím mobile mở khi nhập liệu không che nút đăng nhập", "Mobile keyboard does not hide the login button", "Form vẫn cuộn hoặc co hợp lý để bấm đăng nhập.", "Form scrolls or resizes so the sign-in button remains reachable.", "small phone viewport"),
        ("Khởi động lại app sau khi đăng nhập", "Restart app after login", "App xử lý session nhất quán, không vào sai role hoặc màn hình trắng.", "App handles session consistently without wrong role or blank screen.", "cold start"),
    ])

    add("01_mobile_app/02_dashboard_navigation/mobile_dashboard_navigation_test_cases.xlsx", "MOB-NAV", "Mobile - Dashboard và điều hướng", "Mobile - Dashboard and Navigation", "Kiểm tra dashboard, drawer/sidebar và chuyển màn hình trên mobile", "Dashboard và điều hướng", "Mobile app > Dashboard/Shell", mobile_pre, [
        ("Dashboard tải số liệu tổng quan sau đăng nhập", "Dashboard loads overview metrics after login", "Các chỉ số doanh thu, bàn, đơn và công nợ hiển thị đúng dữ liệu backend.", "Revenue, tables, orders, and debts metrics match backend data."),
        ("Nút refresh trên dashboard cập nhật lại dữ liệu", "Dashboard refresh updates data", "Sau khi refresh, số liệu mới được lấy lại và không nhân đôi dữ liệu.", "After refresh, metrics are reloaded without duplicate data."),
        ("Mở drawer điều hướng trên màn hình nhỏ", "Open navigation drawer on small screen", "Drawer mở mượt, không che lỗi nội dung chính và có đủ menu theo role.", "Drawer opens smoothly and lists role-based menu items."),
        ("Chuyển từ dashboard sang quản lý bàn", "Navigate from dashboard to table management", "Màn hình quản lý bàn mở đúng và có nút quay lại/drawer hợp lệ.", "Table management opens correctly with valid back/drawer behavior."),
        ("Chuyển sang hóa đơn sau khi có đơn mở", "Navigate to invoices after creating an open order", "Danh sách hóa đơn phản ánh đơn vừa tạo.", "Invoice list reflects the newly created order."),
        ("Staff không thấy menu web admin trong điều hướng", "Staff does not see web admin in navigation", "Menu điều hướng của staff chỉ có các module được phép.", "Staff navigation only contains permitted modules.", "role=staff"),
        ("Admin thấy đầy đủ module nhân viên, chấm công, thống kê", "Admin sees employees, attendance, and reporting modules", "Các mục quản trị xuất hiện và chuyển trang được.", "Admin modules are visible and navigable.", "role=admin"),
        ("Đổi hướng portrait sang landscape", "Switch from portrait to landscape", "Bố cục tự co giãn, không tràn chữ hoặc mất nút thao tác.", "Layout adapts without clipped text or lost actions.", "Android rotation"),
        ("Mất mạng khi dashboard đang tải", "Network loss while dashboard is loading", "App hiển thị trạng thái lỗi/empty thay vì crash.", "App shows error/empty state instead of crashing."),
        ("Ảnh thương hiệu hoặc ảnh món ăn lỗi URL", "Brand or menu image URL fails to load", "App hiển thị fallback icon/hình thay thế.", "App displays fallback icon/image."),
        ("Thông tin tiền tệ hiển thị định dạng VNĐ", "Currency values use VND formatting", "Số tiền có dấu phân cách và hậu tố đ phù hợp.", "Money values use separators and VND-style suffix."),
        ("Không xuất hiện chữ mojibake ở các nhãn chính", "No mojibake appears in primary labels", "Tiếng Việt hiển thị đúng dấu ở app bar, menu và nút.", "Vietnamese labels render with correct diacritics."),
    ])

    add("01_mobile_app/03_table_management/mobile_table_management_test_cases.xlsx", "MOB-TABLE", "Mobile - Quản lý bàn", "Mobile - Table Management", "Kiểm tra sơ đồ bàn, nhấn nhanh gọi món và nhấn giữ sửa/xóa bàn", "Quản lý bàn", "Mobile app > Quản lý bàn", mobile_pre, [
        ("Sơ đồ bàn hiển thị đủ bàn từ backend", "Table map displays all backend tables", "Tất cả bàn từ API /tables xuất hiện với tên, khu vực, số ghế và trạng thái.", "All tables from /tables appear with name, area, seats, and status."),
        ("Card bàn không bị bottom overflow trên màn hình 480x864", "Table cards do not bottom-overflow at 480x864", "Không xuất hiện cảnh báo overflow màu vàng/đen; nút vẫn bấm được.", "No yellow/black overflow warning appears and actions remain tappable.", "480x864 viewport"),
        ("Nhấn nhanh vào bàn trống mở màn hình gọi món", "Quick tap available table opens order entry", "App điều hướng sang danh sách món ăn của đúng bàn.", "App navigates to menu/order entry for the selected table."),
        ("Nhấn giữ bàn mở sheet thêm/sửa/xóa", "Long press table opens add/edit/delete sheet", "Bottom sheet có các hành động quản lý bàn hợp lệ.", "Bottom sheet shows valid table management actions."),
        ("Thêm bàn mới với tên, số ghế và khu vực", "Add new table with name, seats, and area", "Bàn mới xuất hiện sau khi lưu và có id không trùng.", "New table appears after saving and has a unique id.", "Bàn Test / 4 / Tầng 2"),
        ("Sửa tên bàn và số ghế", "Edit table name and seats", "Thông tin bàn cập nhật trên grid và trong API.", "Updated table info appears in grid and API."),
        ("Xóa bàn chưa có đơn mở", "Delete table without open order", "Bàn biến mất khỏi danh sách và không ảnh hưởng bàn khác.", "Table is removed without affecting other tables."),
        ("Thử xóa bàn đang có đơn mở", "Attempt to delete table with open order", "Hệ thống chặn hoặc xử lý an toàn để không mất hóa đơn.", "System blocks or safely handles deletion without losing invoice."),
        ("Trạng thái bàn đặt trước hiển thị chip màu cam", "Reserved table shows orange status chip", "Bàn đặt trước có nhãn Đặt trước và màu trạng thái phù hợp.", "Reserved table displays reserved label and matching color."),
        ("Bàn đang dùng đổi trạng thái sau khi tạo đơn", "Occupied status changes after order creation", "Sau khi thêm món, bàn chuyển sang trạng thái đang dùng/occupied.", "After adding items, table becomes occupied."),
        ("Bấm chip hoặc nút Gọi món trong card", "Tap the call-order action chip inside table card", "Nút mở cùng màn hình gọi món như nhấn vào card.", "Action chip opens the same order-entry screen."),
        ("Danh sách nhiều bàn vẫn cuộn mượt", "Large table list scrolls smoothly", "Grid không bị giật, không chồng nút khi có nhiều bàn.", "Grid remains smooth without overlapping actions."),
        ("Tên bàn quá dài được rút gọn hợp lý", "Very long table name is truncated cleanly", "Tên dài không phá vỡ layout, có ellipsis hoặc xuống dòng an toàn.", "Long name does not break layout and is safely truncated/wrapped."),
        ("Nhập số ghế không hợp lệ khi thêm bàn", "Enter invalid seat count when adding a table", "App dùng giá trị mặc định hoặc báo lỗi, không crash.", "App defaults or shows validation error without crashing.", "abc, -1, 0"),
        ("Refresh sau khi backend reset dữ liệu", "Refresh after backend data reset", "Grid tải lại đúng 6 bàn seed ban đầu.", "Grid reloads the six seed tables correctly.", "/reset"),
        ("Bàn VIP hiển thị khu vực riêng và số ghế lớn", "VIP table displays area and larger seat count", "Bàn VIP có tên/khu vực/số ghế đúng và không tràn text.", "VIP table shows correct data without text overflow."),
    ])

    add("01_mobile_app/04_order_entry/mobile_order_entry_test_cases.xlsx", "MOB-ORDER", "Mobile - Gọi món", "Mobile - Order Entry", "Kiểm tra tạo đơn, thêm món bằng dấu cộng và cập nhật số lượng", "Gọi món", "Mobile app > Bàn > Gọi món", mobile_pre, [
        ("Mở bàn chưa có hóa đơn hiển thị trạng thái chưa có hóa đơn", "Open table without invoice shows no-invoice state", "Khối hóa đơn báo chưa có hóa đơn và tổng tiền bằng 0.", "Invoice panel shows no order and total is zero."),
        ("Danh sách món ăn chỉ hiển thị món đang bán", "Menu list shows only available items", "Món unavailable không xuất hiện ở màn hình gọi món.", "Unavailable items are hidden from order entry."),
        ("Bấm dấu cộng ở món đầu tiên tạo hóa đơn mới", "Tap plus on first menu item to create a new invoice", "Đơn mới được tạo, xuất hiện mã hóa đơn và món đã thêm.", "New order is created with order id and selected item.", "Phở bò x1"),
        ("Bấm dấu cộng nhiều lần tăng số lượng món", "Tap plus multiple times to increase item quantity", "Số lượng và tổng tiền tăng đúng theo đơn giá.", "Quantity and total increase according to unit price."),
        ("Bấm nút trừ giảm số lượng món", "Tap minus to decrease item quantity", "Số lượng giảm và tổng tiền cập nhật ngay.", "Quantity decreases and total updates immediately."),
        ("Giảm số lượng về 0 xóa dòng món khỏi hóa đơn", "Decrease quantity to zero removes item row", "Món bị xóa khỏi danh sách, tổng tiền tính lại.", "Item is removed and total is recalculated."),
        ("Thêm nhiều món khác loại vào cùng hóa đơn", "Add multiple different menu items to one order", "Các dòng món tách biệt, tổng tiền bằng tổng từng dòng.", "Line items stay separate and total equals sum of rows."),
        ("Kéo refresh màn hình gọi món", "Pull to refresh order entry screen", "Đơn mở và menu được tải lại không mất dữ liệu.", "Open order and menu reload without losing data."),
        ("Ảnh món ăn lỗi hiển thị icon fallback", "Broken menu image shows fallback icon", "Ô ảnh không làm crash list, icon thay thế xuất hiện.", "Broken image does not crash list and fallback icon appears."),
        ("Backend trả lỗi khi thêm món không tồn tại", "Backend returns error for non-existing menu item", "App hiển thị lỗi và không thêm dòng sai vào hóa đơn.", "App shows error and does not add invalid item.", "menuItemId=9999"),
        ("Bấm back về sơ đồ bàn sau khi thêm món", "Go back to table map after adding item", "Sơ đồ bàn refresh và trạng thái bàn cập nhật.", "Table map refreshes and table status updates."),
        ("Hai thao tác thêm món liên tiếp không tạo hai đơn mở cho một bàn", "Two quick add actions do not create duplicate open orders", "Một bàn chỉ còn một hóa đơn open duy nhất.", "A table keeps only one open order."),
        ("Tổng tiền không âm khi giảm số lượng liên tục", "Total never becomes negative when decreasing repeatedly", "Tổng tiền nhỏ nhất là 0 và không hiển thị NaN.", "Total minimum is zero and never shows NaN."),
        ("Món có giá lớn vẫn hiển thị tiền đúng", "High-priced item displays money correctly", "Định dạng tiền không tràn layout và tính đúng.", "Currency formatting stays correct without overflow."),
        ("Tên món dài trong danh sách món ăn", "Long menu item name in order entry", "Tên món dài không che nút cộng.", "Long item name does not cover the plus button."),
        ("Không có món đang bán", "No available menu items", "App hiển thị empty state hoặc danh sách rỗng rõ ràng.", "App shows clear empty state or empty list."),
    ])

    add("01_mobile_app/05_menu_management/mobile_menu_management_test_cases.xlsx", "MOB-MENU", "Mobile - Quản lý món ăn", "Mobile - Menu Management", "Kiểm tra danh mục món, ảnh món, thêm/sửa/xóa/bật tắt món", "Quản lý món ăn", "Mobile app > Quản lý món ăn", mobile_pre, [
        ("Danh sách món ăn tải từ API /menu", "Menu list loads from /menu API", "Mỗi món có tên, danh mục, giá, trạng thái và ảnh.", "Each item has name, category, price, status, and image."),
        ("Admin thêm món mới đầy đủ thông tin", "Admin creates a complete menu item", "Món mới xuất hiện trong danh sách và có id không trùng.", "New item appears with unique id.", "Bún bò / Món chính / 65000"),
        ("Admin sửa tên, danh mục và giá món", "Admin edits item name, category, and price", "Thông tin món cập nhật và màn hình gọi món nhận dữ liệu mới.", "Item updates and order-entry screen receives new data."),
        ("Admin xóa món chưa dùng trong hóa đơn", "Admin deletes item not used by invoices", "Món bị xóa khỏi quản lý món và không còn ở gọi món.", "Item disappears from menu management and order entry."),
        ("Admin bật/tắt trạng thái đang bán", "Admin toggles availability", "Món tắt bán không xuất hiện ở gọi món nhưng vẫn thấy trong quản lý món.", "Unavailable item is hidden from order entry but visible in menu management."),
        ("Staff chỉ xem danh sách món, không thấy thao tác phá dữ liệu", "Staff views menu without destructive controls", "Staff không được thêm/sửa/xóa/tắt bán nếu chính sách quyền áp dụng.", "Staff cannot add/edit/delete/toggle items when permission policy is enforced.", "role=staff"),
        ("Nhập giá món bằng chữ", "Enter non-numeric menu price", "App báo lỗi hoặc đưa về giá an toàn, không crash.", "App shows validation or safe default without crashing.", "abc"),
        ("Nhập giá món âm", "Enter negative menu price", "Hệ thống không lưu giá âm cho món ăn.", "System does not save a negative item price.", "-10000"),
        ("URL ảnh món hợp lệ tải thành công", "Valid item image URL loads successfully", "Ảnh hiển thị đúng kích thước trong card/list.", "Image displays at expected size in card/list.", "Unsplash URL"),
        ("URL ảnh món sai hiển thị fallback", "Invalid item image URL shows fallback", "Không làm vỡ layout, có icon thay thế.", "Layout remains intact with fallback icon."),
        ("Tên món trùng với món đã có", "Duplicate menu item name", "App/API xử lý rõ ràng theo quy tắc sản phẩm, không tạo dữ liệu khó phân biệt.", "App/API handles duplicates clearly per product rule."),
        ("Danh mục mới được lưu cùng món", "New category is saved with item", "Danh mục hiển thị nhất quán trong quản lý món và gọi món.", "Category is consistent in menu management and order entry."),
        ("Món đang nằm trong hóa đơn cũ vẫn giữ snapshot giá", "Item used in old invoice keeps price snapshot", "Hóa đơn cũ không bị đổi tổng tiền khi sửa giá món.", "Old invoice total does not change after item price edit."),
        ("Danh sách món dài cuộn không giật", "Long menu list scrolls smoothly", "Các ảnh lazy load/fallback không gây nhảy layout lớn.", "Images lazy load/fallback without large layout shifts."),
    ])

    add("01_mobile_app/06_invoice_debt/mobile_invoice_debt_test_cases.xlsx", "MOB-BILL", "Mobile - Hóa đơn và công nợ", "Mobile - Invoices and Debts", "Kiểm tra hóa đơn, thanh toán, ghi nợ và thu nợ trên mobile", "Hóa đơn và công nợ", "Mobile app > Hóa đơn/Công nợ", mobile_pre, [
        ("Danh sách hóa đơn hiển thị đơn open mới tạo", "Invoice list displays newly opened order", "Hóa đơn có bàn, mã, trạng thái, nhân viên và tổng tiền.", "Invoice shows table, id, status, staff, and total."),
        ("Mở rộng hóa đơn hiển thị chi tiết món", "Expand invoice to show line items", "Mỗi dòng món có tên, số lượng và thành tiền.", "Each line item has name, quantity, and line total."),
        ("Admin bấm Đã trả cho hóa đơn open", "Admin marks an open invoice as paid", "Hóa đơn đổi sang paid, bàn được giải phóng và thống kê doanh thu tăng.", "Invoice becomes paid, table is released, and revenue increases."),
        ("Staff được phép bấm Đã trả theo quyền thu tiền", "Staff can mark paid when collection permission allows", "Staff có nút Đã trả và thanh toán thành công nếu canCollect true.", "Staff has paid button and payment succeeds when canCollect is true.", "role=staff"),
        ("Vai trò không có quyền không thấy nút Đã trả", "Unauthorized role does not see paid button", "Nút thanh toán bị ẩn hoặc API từ chối quyền.", "Payment button is hidden or API rejects permission."),
        ("Nhập giảm giá khi thanh toán", "Enter discount when paying invoice", "Tổng thanh toán áp dụng giảm giá hợp lệ, không âm.", "Payable amount applies valid discount and never goes negative.", "discount=10000"),
        ("Nhập giảm giá lớn hơn tổng tiền", "Enter discount greater than total", "Hệ thống chặn hoặc đưa tổng về 0, không âm.", "System blocks or clamps payable amount to zero.", "discount > total"),
        ("Ghi công nợ từ hóa đơn open", "Convert open invoice to debt", "Tạo bản ghi công nợ có khách hàng, số điện thoại, ghi chú và tổng tiền.", "Debt record is created with customer, phone, note, and amount."),
        ("Bỏ trống tên khách khi ghi nợ", "Leave customer name empty when creating debt", "App/API báo lỗi hoặc lưu nhãn khách vãng lai theo quy tắc.", "App/API validates or stores walk-in customer per product rule."),
        ("Danh sách công nợ hiển thị khoản nợ mới", "Debt list displays newly created debt", "Khoản nợ có trạng thái unpaid và liên kết hóa đơn gốc.", "Debt has unpaid status and source invoice link."),
        ("Thu nợ toàn bộ khoản công nợ", "Collect full debt amount", "Khoản nợ chuyển paid/collected và không còn trong danh sách nợ mở.", "Debt becomes paid/collected and leaves open debt list."),
        ("Thu nợ một phần nếu API hỗ trợ", "Collect partial debt if API supports it", "Số tiền còn lại tính đúng hoặc hệ thống từ chối rõ ràng nếu chưa hỗ trợ.", "Remaining amount is correct or system clearly rejects if unsupported."),
        ("Refresh hóa đơn sau khi thanh toán", "Refresh invoices after payment", "Không hiển thị lại nút thanh toán cho hóa đơn đã paid.", "Paid invoice no longer shows pay button."),
        ("Hóa đơn rỗng không thanh toán được", "Empty invoice cannot be paid", "API chặn hóa đơn không có item hoặc tổng bằng 0 theo quy tắc.", "API blocks invoices with no items or zero total per rule."),
        ("Công nợ có số điện thoại sai định dạng", "Debt with invalid phone format", "App cảnh báo hoặc lưu theo chính sách, không crash.", "App warns or stores per policy without crashing.", "abc123"),
        ("Tiền hóa đơn/công nợ giữ định dạng VNĐ", "Invoice/debt amounts retain VND formatting", "Mọi số tiền ở danh sách và chi tiết hiển thị nhất quán.", "All list/detail amounts display consistently."),
    ])

    add("01_mobile_app/07_statistics/mobile_statistics_test_cases.xlsx", "MOB-STATS", "Mobile - Thống kê", "Mobile - Statistics", "Kiểm tra thống kê theo ngày, tháng, năm và liên kết hóa đơn", "Thống kê", "Mobile app > Thống kê", mobile_pre, [
        ("Mở thống kê ngày hiện tại", "Open current-day statistics", "Doanh thu, số hóa đơn và công nợ phản ánh dữ liệu trong ngày.", "Revenue, invoices, and debts reflect current-day data."),
        ("Chọn thống kê theo tháng", "Select monthly statistics", "Số liệu tháng bao gồm các hóa đơn paid trong tháng.", "Monthly metrics include paid invoices in the month."),
        ("Chọn thống kê theo năm", "Select yearly statistics", "Số liệu năm tổng hợp đúng theo tháng/quý nếu có.", "Yearly metrics aggregate correctly by month/quarter if available."),
        ("Không tính hóa đơn open vào doanh thu", "Open invoices are excluded from revenue", "Doanh thu chỉ tính hóa đơn đã trả.", "Revenue only includes paid invoices."),
        ("Ghi nợ không tính như doanh thu đã thu", "Debt is not counted as collected revenue", "Công nợ được hiển thị riêng với doanh thu.", "Debts are reported separately from collected revenue."),
        ("Thanh toán hóa đơn rồi refresh thống kê", "Pay invoice then refresh statistics", "Doanh thu và số hóa đơn tăng đúng một lần.", "Revenue and invoice count increase exactly once."),
        ("Reset dữ liệu demo rồi mở thống kê", "Reset demo data then open statistics", "Thống kê về 0 hoặc seed mặc định rõ ràng.", "Statistics return to zero or known seed state."),
        ("Ngày không có giao dịch", "Date range with no transactions", "Hiển thị empty state/0 thay vì lỗi.", "Shows empty state/zero instead of error."),
        ("Giá trị doanh thu lớn", "Large revenue values", "Số tiền lớn không tràn layout và format đúng.", "Large amounts do not overflow and are formatted correctly."),
        ("Mất kết nối khi tải thống kê", "Connection loss while loading statistics", "App hiển thị lỗi có thể thử lại.", "App shows retryable error."),
        ("Dữ liệu thống kê khớp danh sách hóa đơn", "Statistics match invoice list", "Tổng doanh thu bằng tổng các hóa đơn paid trong kỳ.", "Revenue equals sum of paid invoices in the period."),
        ("Đổi timezone thiết bị", "Change device timezone", "Ngày thống kê vẫn nhất quán với dữ liệu backend.", "Report date remains consistent with backend data."),
    ])

    add("01_mobile_app/08_employee_attendance/mobile_employee_attendance_test_cases.xlsx", "MOB-HR", "Mobile - Nhân viên và chấm công", "Mobile - Employees and Attendance", "Kiểm tra danh sách nhân viên, chấm công vào/ra và quyền admin", "Nhân viên và chấm công", "Mobile app > Nhân viên/Chấm công", mobile_pre, [
        ("Admin xem danh sách nhân viên", "Admin views employee list", "Danh sách có tên, vai trò, số điện thoại, lương ngày và trạng thái active.", "List shows name, role, phone, daily salary, and active status.", "role=admin"),
        ("Admin thêm nhân viên mới", "Admin creates a new employee", "Nhân viên mới xuất hiện với id không trùng và active mặc định hợp lệ.", "New employee appears with unique id and valid default active state."),
        ("Admin sửa thông tin nhân viên", "Admin edits employee information", "Thông tin cập nhật ở danh sách và liên kết payroll/chấm công.", "Updates appear in list and linked payroll/attendance."),
        ("Admin khóa nhân viên không còn hoạt động", "Admin deactivates an employee", "Nhân viên inactive không được tính công/lương mới nếu chính sách áp dụng.", "Inactive employee is excluded from new attendance/payroll if policy applies."),
        ("Staff không truy cập màn quản lý nhân viên", "Staff cannot access employee management", "Menu bị ẩn hoặc màn hình từ chối quyền.", "Menu is hidden or screen denies access.", "role=staff"),
        ("Nhân viên check-in đầu ngày", "Employee checks in at start of shift", "Bản ghi chấm công tạo với thời gian check-in và trạng thái đang làm.", "Attendance record has check-in time and working status."),
        ("Nhân viên check-out cuối ngày", "Employee checks out at end of shift", "Bản ghi cập nhật check-out và tổng thời gian làm.", "Record updates with check-out and worked time."),
        ("Check-in hai lần trong cùng ngày", "Check in twice on same day", "Hệ thống chặn hoặc cập nhật an toàn, không tạo công trùng.", "System blocks or safely handles duplicate attendance."),
        ("Check-out khi chưa check-in", "Check out without prior check-in", "App/API báo lỗi hợp lệ, không tạo record sai.", "App/API shows valid error without invalid record."),
        ("Danh sách chấm công lọc theo ngày", "Attendance list filters by date", "Chỉ hiển thị record đúng ngày chọn.", "Only records for selected date are shown."),
        ("Nhân viên nghỉ không có record trong ngày", "Absent employee has no record for the day", "Không tự tạo công hoặc trạng thái được hiển thị rõ.", "No automatic attendance or clear absence state."),
        ("Chấm công khi mất mạng", "Attendance action while offline", "App báo lỗi và không lưu nửa chừng.", "App shows error and does not partially save."),
        ("Tên nhân viên dài trong danh sách", "Long employee name in list", "Không làm tràn card/list item.", "Does not overflow card/list item."),
        ("Số điện thoại nhân viên sai định dạng", "Invalid employee phone format", "Form xử lý theo quy tắc validation, không crash.", "Form handles validation per rules without crashing."),
    ])

    add("01_mobile_app/09_notification_payroll/mobile_notification_payroll_test_cases.xlsx", "MOB-OPS", "Mobile - Thông báo và tính lương", "Mobile - Notifications and Payroll", "Kiểm tra thông báo nội bộ và tính tiền nhân viên", "Thông báo và tính lương", "Mobile app > Thông báo/Tính lương", mobile_pre, [
        ("Tải danh sách thông báo cho tất cả nhân viên", "Load notifications for all employees", "Thông báo audience all hiển thị cho admin và staff.", "Audience all notifications are visible to admin and staff."),
        ("Admin tạo thông báo mới", "Admin creates a new notification", "Thông báo mới có tiêu đề, nội dung, audience và thời gian tạo.", "New notification has title, body, audience, and created time."),
        ("Staff đọc thông báo nhưng không tạo thông báo quản trị", "Staff reads notifications but cannot create admin notices", "Staff không thấy nút tạo hoặc API từ chối quyền.", "Staff lacks create action or API rejects permission."),
        ("Thông báo tiêu đề trống", "Notification with empty title", "Form/API báo lỗi hợp lệ.", "Form/API returns validation error."),
        ("Nội dung thông báo dài", "Long notification body", "Nội dung xuống dòng đúng và không phá layout.", "Long body wraps without breaking layout."),
        ("Mở màn hình tính lương theo tháng", "Open monthly payroll screen", "Danh sách lương tính theo nhân viên active và ngày công.", "Payroll is calculated from active employees and attendance days."),
        ("Lương nhân viên tăng theo số ngày chấm công", "Employee salary increases with attendance days", "Tổng lương bằng lương ngày nhân số công hợp lệ.", "Total salary equals daily salary multiplied by valid work days."),
        ("Nhân viên không chấm công có lương 0 trong kỳ", "Employee without attendance has zero salary for period", "Payroll không tự cộng lương cho người không có record.", "Payroll does not add salary without records."),
        ("Admin lọc bảng lương theo tháng khác", "Admin filters payroll by another month", "Số ngày công và lương thay đổi theo kỳ chọn.", "Work days and salary change by selected period."),
        ("Staff không xem được bảng lương toàn bộ nếu bị hạn quyền", "Staff cannot view full payroll when restricted", "Menu bị ẩn hoặc chỉ xem dữ liệu của bản thân theo chính sách.", "Menu hidden or staff can only view own data per policy."),
        ("Payroll xử lý lương ngày bằng 0", "Payroll handles zero daily salary", "Tổng lương hiển thị 0 và không lỗi chia/tính toán.", "Total salary is zero without calculation errors."),
        ("Refresh thông báo và payroll sau khi thay đổi dữ liệu", "Refresh notifications and payroll after data changes", "Dữ liệu mới xuất hiện đúng một lần, không nhân đôi dòng.", "Updated data appears once without duplicate rows."),
    ])

    add("02_web_app/01_authentication_layout/web_authentication_layout_test_cases.xlsx", "WEB-AUTH", "Web - Đăng nhập và bố cục", "Web - Authentication and Layout", "Kiểm tra đăng nhập web, responsive và thương hiệu", "Đăng nhập web", "Web app > Login page", web_pre, [
        ("Mở web hiển thị tên app Quản Lý Nhà Hàng", "Open web and display Restaurant Management app name", "Tên app, logo và form đăng nhập hiển thị đúng trên trình duyệt.", "App name, logo, and login form display correctly."),
        ("Đăng nhập admin trên Chrome desktop", "Admin login on Chrome desktop", "Web vào dashboard với quyền admin.", "Web opens dashboard with admin privileges.", "admin / duyencocong2"),
        ("Đăng nhập staff trên Chrome desktop", "Staff login on Chrome desktop", "Web vào dashboard với quyền staff.", "Web opens dashboard with staff privileges.", "staff / duyencocong2"),
        ("Đăng nhập sai mật khẩu trên web", "Wrong password on web login", "Hiển thị lỗi rõ ràng, không chuyển trang.", "Shows clear error and does not navigate.", "wrong password"),
        ("Reload trình duyệt ở màn login", "Reload browser on login page", "Trang vẫn tải lại sạch, không lỗi trắng.", "Page reloads cleanly without blank screen."),
        ("Bố cục login ở 1366x768", "Login layout at 1366x768", "Form cân đối, hình ảnh không bị cắt xấu.", "Form is balanced and images are not badly cropped."),
        ("Bố cục login ở 390x844", "Login layout at 390x844", "Form vừa màn hình mobile web và không tràn ngang.", "Form fits mobile web without horizontal overflow."),
        ("Mất backend khi web đăng nhập", "Backend unavailable during web login", "Web báo lỗi kết nối, không treo spinner.", "Web shows connection error and spinner stops."),
        ("Ảnh nền/ảnh hero không tải được", "Hero/background image fails to load", "Web vẫn đọc được form và dùng màu/fallback hợp lý.", "Form remains readable with fallback styling."),
        ("Nhập liệu bằng phím Enter", "Submit login with Enter key", "Form đăng nhập xử lý giống bấm nút.", "Login behaves the same as button submit."),
    ])

    add("02_web_app/02_sidebar_navigation/web_sidebar_navigation_test_cases.xlsx", "WEB-NAV", "Web - Sidebar và điều hướng", "Web - Sidebar and Navigation", "Kiểm tra sidebar desktop, drawer tablet/mobile và module web mở rộng", "Sidebar web", "Web app > Shell/Sidebar", web_pre, [
        ("Desktop hiển thị sidebar cố định", "Desktop displays persistent sidebar", "Sidebar nằm trái, nội dung chính không bị che.", "Left sidebar is persistent and does not cover main content."),
        ("Tablet/mobile chuyển sang drawer", "Tablet/mobile switches to drawer", "Drawer mở bằng icon menu và đóng đúng.", "Drawer opens with menu icon and closes correctly."),
        ("Chọn từng menu đổi màn hình không reload toàn trang", "Selecting menu changes screen without full page reload", "Nội dung chuyển mượt và giữ session.", "Content switches smoothly while keeping session."),
        ("Admin thấy đủ module web mở rộng", "Admin sees full extended web modules", "Có quản lý nhân viên, thông báo, tính lương, web admin.", "Employees, notifications, payroll, and web admin are available."),
        ("Staff không thấy module admin web", "Staff does not see web admin modules", "Sidebar staff chỉ hiển thị chức năng được phép.", "Staff sidebar only shows permitted features."),
        ("Mục menu active được tô nổi bật", "Active menu item is highlighted", "Mục đang mở có trạng thái active rõ ràng.", "Current menu has a clear active state."),
        ("Refresh browser khi đang ở một module", "Refresh browser while inside a module", "Web không mất layout hoặc rơi vào màn trắng.", "Web does not lose layout or show blank screen."),
        ("Logout từ web sidebar", "Logout from web sidebar", "Session xóa và quay về login.", "Session is cleared and login page opens."),
        ("Sidebar không tràn khi tên module dài", "Sidebar does not overflow with long labels", "Label được rút gọn/xuống dòng hợp lý.", "Labels truncate/wrap cleanly."),
        ("Điều hướng bằng bàn phím tab", "Keyboard tab navigation", "Focus di chuyển qua menu và nút chính theo thứ tự hợp lý.", "Focus moves through menu and primary actions logically."),
        ("Màu đen trắng xanh nước biển nhất quán", "Black-white-blue palette is consistent", "Các vùng chính dùng palette thương hiệu, không lạc màu.", "Primary surfaces use the brand palette consistently."),
        ("Không có nested card gây rối ở layout web", "No confusing nested cards in web layout", "Các khu vực chính rõ ràng, card chỉ dùng cho item/tool.", "Main areas are clear and cards are used only for items/tools."),
    ])

    add("02_web_app/03_table_order_web/web_table_order_test_cases.xlsx", "WEB-TABLE", "Web - Bàn và gọi món", "Web - Tables and Orders", "Kiểm tra quản lý bàn và gọi món trên website", "Bàn và gọi món web", "Web app > Quản lý bàn/Gọi món", web_pre, [
        ("Grid bàn web dùng nhiều cột trên desktop", "Web table grid uses multiple desktop columns", "Bàn tận dụng chiều rộng desktop, không quá thưa hoặc quá hẹp.", "Tables use desktop width without being too sparse or cramped."),
        ("Nhấn bàn trên web mở gọi món đúng bàn", "Click table on web opens order entry for correct table", "Tên bàn trong app bar/order entry khớp bàn đã chọn.", "Table name in order entry matches selected table."),
        ("Nhấn giữ hoặc thao tác tương đương để sửa/xóa bàn trên web", "Long press or equivalent action edits/deletes table on web", "Web có cách thao tác phù hợp chuột/touch để quản lý bàn.", "Web provides mouse/touch-friendly table management."),
        ("Thêm bàn trên desktop web", "Add table on desktop web", "Dialog/form không bị che, lưu xong grid refresh.", "Dialog/form is visible and grid refreshes after save."),
        ("Sửa bàn trên web", "Edit table on web", "Thông tin cập nhật ngay trong grid.", "Updated table info appears immediately in grid."),
        ("Xóa bàn trên web", "Delete table on web", "Bàn bị xóa và các chỉ số dashboard cập nhật.", "Table is deleted and dashboard metrics update."),
        ("Thêm món bằng dấu cộng trên web", "Add item with plus button on web", "Hóa đơn tạo/cập nhật giống mobile.", "Invoice is created/updated like mobile."),
        ("Danh sách món web có ảnh và fallback", "Web menu list has images and fallback", "Ảnh tải đẹp, lỗi ảnh không vỡ layout.", "Images render nicely and failures do not break layout."),
        ("Số lượng món cập nhật liên tục trên web", "Item quantity updates repeatedly on web", "Không double-submit ngoài ý muốn, tổng tiền đúng.", "No unintended double-submit and total is correct."),
        ("Web refresh sau khi tạo hóa đơn", "Web refresh after creating invoice", "Đơn open vẫn còn và bàn đúng trạng thái.", "Open order persists and table status is correct."),
        ("Hai tab trình duyệt cùng thao tác một bàn", "Two browser tabs operate on same table", "Dữ liệu cuối cùng nhất quán, không tạo đơn trùng.", "Final data stays consistent without duplicate open orders."),
        ("Bố cục gọi món ở màn hình hẹp", "Order-entry layout on narrow browser width", "Nút cộng/trừ và tổng tiền không chồng nhau.", "Plus/minus buttons and total do not overlap."),
        ("Tốc độ tải danh sách nhiều món", "Performance with many menu items", "Danh sách vẫn cuộn và thao tác được trong ngưỡng chấp nhận.", "List remains scrollable and usable within acceptable performance."),
        ("Thông báo lỗi API khi lưu bàn web", "API error notification when saving table on web", "Web hiển thị snackbar/dialog lỗi và giữ dữ liệu người nhập.", "Web shows error and preserves user input."),
    ])

    add("02_web_app/04_admin_console/web_admin_console_test_cases.xlsx", "WEB-ADMIN", "Web - Admin console", "Web - Admin Console", "Kiểm tra trang web admin nhiều chức năng hơn app mobile", "Web admin", "Web app > Web Admin", web_pre, [
        ("Admin mở trang Web Admin", "Admin opens Web Admin page", "Trang hiển thị tình trạng hệ thống, dữ liệu và công cụ quản trị.", "Page shows system status, data summary, and admin tools.", "role=admin"),
        ("Staff không thấy Web Admin", "Staff cannot see Web Admin", "Menu web admin bị ẩn hoặc truy cập bị chặn.", "Web Admin menu is hidden or access is denied.", "role=staff"),
        ("Kiểm tra health backend từ web admin", "Check backend health from web admin", "Trạng thái ok, số bàn, món, nhân viên hiển thị đúng.", "Status ok plus table/menu/employee counts are displayed."),
        ("Xuất dữ liệu JSON từ web admin", "Export JSON data from web admin", "Dữ liệu export có users an toàn, tables, menu, orders, debts, employees.", "Export includes safe users, tables, menu, orders, debts, employees."),
        ("Reset demo data từ web admin", "Reset demo data from web admin", "Dữ liệu quay về seed, dashboard/table/menu refresh theo seed.", "Data returns to seed and dashboard/tables/menu refresh."),
        ("Không cho reset nếu chưa xác nhận", "Reset is blocked without confirmation", "Hành động nguy hiểm cần xác nhận rõ.", "Dangerous action requires explicit confirmation."),
        ("Hiển thị checklist vận hành nhà hàng", "Display restaurant operation checklist", "Checklist có các mục cần kiểm trước ca/ngày.", "Checklist contains pre-shift/daily operation items."),
        ("Thống kê tháng trong web admin", "Monthly stats in web admin", "Web admin hiển thị thống kê tháng từ API stats.", "Web admin displays monthly stats from stats API."),
        ("Bảng lương trong web admin", "Payroll summary in web admin", "Admin xem được tổng lương và ngày công.", "Admin can see salary and attendance summary."),
        ("Lỗi admin/health hiển thị rõ", "admin/health error is shown clearly", "Không crash toàn trang; có trạng thái lỗi.", "Page does not crash and shows error state."),
        ("Export dữ liệu lớn", "Export large data set", "Preview JSON vẫn cuộn/hiển thị được, không khóa trình duyệt lâu.", "JSON preview remains scrollable without freezing browser."),
        ("Web admin không lộ password trong export", "Web admin export does not expose passwords", "Dữ liệu export không chứa password plain text nếu chính sách bảo mật áp dụng.", "Export does not include plain-text passwords when security policy applies."),
        ("Admin console dùng icon/nút rõ nghĩa", "Admin console uses clear icons/buttons", "Các nút health/export/reset có icon và tooltip/label rõ.", "Health/export/reset buttons have clear icons and labels/tooltips."),
        ("Responsive web admin trên tablet", "Responsive admin console on tablet", "Các khối xếp lại hợp lý, không tràn ngang.", "Panels reflow without horizontal overflow."),
        ("Admin console đồng bộ sau khi mobile đổi dữ liệu", "Admin console syncs after mobile data changes", "Refresh web admin thấy dữ liệu mobile vừa tạo.", "Refreshing admin console shows data created from mobile."),
        ("Ghi log hoặc hiển thị thời gian generatedAt", "Show generatedAt/log time", "Thời gian health/export giúp biết dữ liệu mới hay cũ.", "Health/export generated time indicates data freshness."),
    ])

    add("02_web_app/05_reporting_export/web_reporting_export_test_cases.xlsx", "WEB-REPORT", "Web - Báo cáo và xuất dữ liệu", "Web - Reports and Export", "Kiểm tra thống kê web theo ngày/tháng/năm và xuất dữ liệu", "Báo cáo web", "Web app > Thống kê/Web Admin Export", web_pre, [
        ("Web thống kê theo ngày", "Web daily report", "Số liệu theo ngày khớp API /stats?period=day.", "Daily metrics match /stats?period=day."),
        ("Web thống kê theo tháng", "Web monthly report", "Số liệu theo tháng khớp API /stats?period=month.", "Monthly metrics match /stats?period=month."),
        ("Web thống kê theo năm", "Web yearly report", "Số liệu theo năm khớp API /stats?period=year.", "Yearly metrics match /stats?period=year."),
        ("Báo cáo không tính hóa đơn open", "Reports exclude open invoices", "Chỉ hóa đơn paid/collected được tính doanh thu.", "Only paid/collected invoices count as revenue."),
        ("Báo cáo hiển thị công nợ riêng", "Reports show debts separately", "Công nợ không bị trộn vào doanh thu đã thu.", "Debts are not mixed with collected revenue."),
        ("Export JSON sau khi thanh toán", "Export JSON after payment", "File/preview chứa order status paid và totals đúng.", "Export/preview contains paid order status and correct totals."),
        ("Export JSON sau khi ghi nợ", "Export JSON after debt creation", "Có bản ghi debts liên kết order đúng.", "Debt record is linked to correct order."),
        ("Dữ liệu xuất không trùng id", "Exported data has no duplicate ids", "Các collection giữ id duy nhất trong từng loại dữ liệu.", "Collections keep unique ids per entity type."),
        ("Refresh báo cáo không nhân đôi kết quả", "Refreshing reports does not duplicate results", "Số liệu ổn định sau nhiều lần refresh.", "Metrics remain stable after repeated refresh."),
        ("Hiển thị số 0 khi chưa có giao dịch", "Show zero when no transactions exist", "Không để ô trống gây hiểu nhầm.", "Shows zero instead of misleading blank cells."),
        ("Bộ lọc ngày/tháng/năm vẫn giữ lựa chọn sau refresh", "Period filter keeps selection after refresh", "Kỳ đang chọn không bị nhảy bất ngờ.", "Selected period does not change unexpectedly."),
        ("In/Copy dữ liệu báo cáo từ trình duyệt", "Print/copy report data from browser", "Nội dung quan trọng đọc được khi in/copy.", "Key report data remains readable when printed/copied."),
    ])

    add("02_web_app/06_browser_responsive/web_browser_responsive_test_cases.xlsx", "WEB-RESP", "Web - Responsive và trình duyệt", "Web - Responsive and Browser Compatibility", "Kiểm tra Chrome/Edge/mobile web, kích thước màn hình và khả năng truy cập", "Responsive web", "Web app > Browser viewports", web_pre, [
        ("Chrome desktop 1366x768", "Chrome desktop 1366x768", "Không tràn ngang, sidebar và nội dung cân đối.", "No horizontal overflow; sidebar and content are balanced.", "Chrome 1366x768"),
        ("Edge desktop 1920x1080", "Edge desktop 1920x1080", "Không kéo giãn card quá mức; nội dung dễ scan.", "Cards do not stretch awkwardly and content is scannable.", "Edge 1920x1080"),
        ("Mobile browser 390x844", "Mobile browser 390x844", "Drawer, bảng bàn và form vẫn thao tác được.", "Drawer, table grid, and forms remain usable.", "390x844"),
        ("Tablet 768x1024", "Tablet 768x1024", "Layout chuyển breakpoint hợp lý.", "Layout switches breakpoints cleanly.", "768x1024"),
        ("Phóng to trình duyệt 125%", "Browser zoom 125%", "Text không chồng lấp nút quan trọng.", "Text does not overlap key buttons.", "zoom=125%"),
        ("Phóng to trình duyệt 150%", "Browser zoom 150%", "Nội dung vẫn cuộn đọc được.", "Content remains scrollable and readable.", "zoom=150%"),
        ("Tab focus qua form và nút", "Tab focus through forms and buttons", "Focus ring rõ, thứ tự tab hợp lý.", "Focus ring is visible and tab order is logical."),
        ("Enter/Space kích hoạt nút chính", "Enter/Space activates primary controls", "Nút có thể dùng bằng bàn phím.", "Buttons can be used by keyboard."),
        ("Độ tương phản màu nút chính", "Primary button color contrast", "Text trắng/xanh/đen đủ tương phản để đọc.", "White/blue/black text has readable contrast."),
        ("Không có scroll ngang toàn trang", "No page-level horizontal scrolling", "Body không xuất hiện scrollbar ngang ở viewport chuẩn.", "Body has no horizontal scrollbar at standard viewports."),
        ("Web app refresh sâu sau build release", "Deep refresh after release build", "Trang build/web phục vụ lại app thay vì lỗi asset.", "Release web build serves app without asset errors."),
        ("Ảnh mạng tải chậm", "Slow network images", "Skeleton/fallback giữ layout ổn định.", "Skeleton/fallback keeps layout stable."),
    ])

    add("03_api_backend/01_auth_api/api_auth_test_cases.xlsx", "API-AUTH", "API - Đăng nhập", "API - Authentication", "Kiểm tra API auth/login và dữ liệu user an toàn", "API đăng nhập", "Backend API > /auth/login", api_pre, [
        ("POST /auth/login với admin hợp lệ", "POST /auth/login with valid admin", "Trả success true, user role admin, không trả password.", "Returns success true, admin role, and no password.", "admin / duyencocong2"),
        ("POST /auth/login với staff hợp lệ", "POST /auth/login with valid staff", "Trả success true, user role staff, không trả password.", "Returns success true, staff role, and no password.", "staff / duyencocong2"),
        ("POST /auth/login sai mật khẩu", "POST /auth/login wrong password", "Trả lỗi 401/400 phù hợp và không trả user.", "Returns proper 401/400 error and no user.", "admin / wrong"),
        ("POST /auth/login thiếu username", "POST /auth/login missing username", "Không crash server, trả lỗi validation.", "Server does not crash and returns validation error."),
        ("POST /auth/login thiếu password", "POST /auth/login missing password", "Không crash server, trả lỗi validation.", "Server does not crash and returns validation error."),
        ("POST /auth/login body JSON sai định dạng", "POST /auth/login malformed JSON", "Express trả lỗi an toàn, process không dừng.", "Express returns safe error and process continues."),
        ("Username phân biệt khoảng trắng đầu/cuối", "Username trims leading/trailing spaces", "Nếu app trim input, login vẫn theo dữ liệu sạch; API không lưu session sai.", "If client trims input, login uses clean data and API does not create wrong session.", " admin "),
        ("Không lộ password trong mọi phản hồi user", "User responses never expose password", "Response user không có trường password plain text.", "User response has no plain-text password field."),
        ("Role legacy được normalize", "Legacy role is normalized", "User cũ được đưa về admin/staff an toàn.", "Legacy user role is normalized to admin/staff safely."),
        ("API root / trả trạng thái server chạy", "API root / returns running status", "Trả success true và message server đang chạy.", "Returns success true and running message.", "GET /"),
    ])

    add("03_api_backend/02_table_menu_api/api_table_menu_test_cases.xlsx", "API-CATALOG", "API - Bàn và món ăn", "API - Tables and Menu", "Kiểm tra CRUD bàn, CRUD món và dữ liệu catalog", "API bàn/món", "Backend API > /tables, /menu", api_pre, [
        ("GET /tables trả danh sách bàn", "GET /tables returns table list", "Mảng bàn có id, name, seats, area, status.", "Table array has id, name, seats, area, status."),
        ("POST /tables thêm bàn hợp lệ", "POST /tables creates valid table", "Bàn mới có id tăng và dữ liệu lưu vào JSON.", "New table has incremented id and is saved to JSON."),
        ("PATCH /tables/:id sửa bàn", "PATCH /tables/:id updates table", "Bàn đúng id được cập nhật, bàn khác không đổi.", "Target table updates and other tables remain unchanged."),
        ("DELETE /tables/:id xóa bàn", "DELETE /tables/:id deletes table", "Bàn không còn trong GET /tables sau khi xóa.", "Table no longer appears in GET /tables."),
        ("DELETE /tables id không tồn tại", "DELETE /tables non-existing id", "Trả 404/lỗi phù hợp, dữ liệu không đổi.", "Returns 404/proper error and data is unchanged.", "id=9999"),
        ("POST /tables thiếu name", "POST /tables missing name", "API xử lý validation hoặc default rõ ràng.", "API validates or defaults clearly."),
        ("POST /tables seats không phải số", "POST /tables non-numeric seats", "Không lưu seats NaN; trả lỗi hoặc default.", "Does not save NaN seats; returns error or default.", "seats=abc"),
        ("GET /menu trả danh sách món", "GET /menu returns menu list", "Mảng món có id, name, category, price, available, imageUrl.", "Menu array has id, name, category, price, available, imageUrl."),
        ("POST /menu thêm món hợp lệ", "POST /menu creates valid item", "Món mới có id tăng và available hợp lệ.", "New item has incremented id and valid availability."),
        ("PATCH /menu/:id sửa giá món", "PATCH /menu/:id updates price", "Giá mới là số hợp lệ và lưu vào JSON.", "New price is numeric and saved."),
        ("PATCH /menu/:id bật/tắt available", "PATCH /menu/:id toggles availability", "Trạng thái available thay đổi và GET /menu phản ánh.", "Availability changes and GET /menu reflects it."),
        ("DELETE /menu/:id xóa món", "DELETE /menu/:id deletes item", "Món không còn trong danh sách quản lý.", "Item no longer appears in management list."),
        ("POST /menu thiếu price", "POST /menu missing price", "API không lưu dữ liệu thiếu giá gây lỗi tính tiền.", "API does not save missing price that breaks totals."),
        ("POST /menu price âm", "POST /menu negative price", "API chặn hoặc chuẩn hóa giá âm.", "API blocks or normalizes negative price.", "-5000"),
        ("Image URL được normalize khi thiếu", "Image URL is normalized when missing", "Món có imageUrl fallback/seed hợp lệ.", "Item has fallback/seed imageUrl."),
        ("nextIds không bị trùng sau reset", "nextIds do not duplicate after reset", "Sau reset và thêm mới, id không trùng dữ liệu seed.", "After reset and creation, id does not collide with seed data."),
    ])

    add("03_api_backend/03_order_invoice_api/api_order_invoice_test_cases.xlsx", "API-ORDER", "API - Đơn và hóa đơn", "API - Orders and Invoices", "Kiểm tra tạo đơn, item, thanh toán và ghi nợ", "API đơn/hóa đơn", "Backend API > /orders", api_pre, [
        ("GET /orders trả danh sách hóa đơn", "GET /orders returns invoices", "Trả mảng đơn với status, items và total.", "Returns orders with status, items, and total."),
        ("GET /orders/table/:id/open khi chưa có đơn", "GET open order for table without order", "Trả null/empty đúng quy ước API.", "Returns null/empty per API contract."),
        ("POST /orders tạo đơn mới cho bàn", "POST /orders creates order for table", "Đơn có id, tableId, staffId, status open và total 0.", "Order has id, tableId, staffId, open status, and zero total."),
        ("POST /orders không tạo trùng đơn open cùng bàn", "POST /orders does not duplicate open order for same table", "Một bàn chỉ có một order open.", "A table has only one open order."),
        ("POST /orders/:id/items thêm món", "POST /orders/:id/items adds item", "Item vào order, total tính đúng.", "Item is added and total is correct."),
        ("POST /orders/:id/items món không tồn tại", "POST item with missing menu id", "Trả lỗi phù hợp, order không đổi.", "Returns error and order remains unchanged.", "menuItemId=9999"),
        ("PATCH /orders/:id/items/:menuItemId tăng số lượng", "PATCH item quantity increases quantity", "Quantity và total tăng đúng.", "Quantity and total increase correctly."),
        ("PATCH /orders/:id/items/:menuItemId về 0", "PATCH item quantity to zero", "Item bị xóa hoặc quantity không âm theo quy tắc.", "Item is removed or quantity never goes negative."),
        ("PATCH item trong order không tồn tại", "PATCH item in non-existing order", "Trả 404/lỗi phù hợp.", "Returns 404/proper error."),
        ("POST /orders/:id/pay bằng role admin", "POST pay with admin role", "Order chuyển paid, paidAt/userName/discount lưu đúng.", "Order becomes paid and payment metadata is saved.", "role=admin"),
        ("POST /orders/:id/pay bằng role staff", "POST pay with staff role", "Staff được phép thanh toán nếu canCollect true.", "Staff can pay when canCollect is true.", "role=staff"),
        ("POST /orders/:id/pay role không hợp lệ", "POST pay with invalid role", "API từ chối quyền thanh toán.", "API rejects unauthorized payment role.", "role=guest"),
        ("POST /orders/:id/pay discount âm", "POST pay with negative discount", "Discount không làm tăng tổng tiền sai hoặc tạo giá trị âm.", "Discount does not create invalid negative/incorrect totals.", "discount=-1000"),
        ("POST /orders/:id/debt tạo công nợ", "POST debt creates debt from order", "Order chuyển debt/unpaid và bản ghi debt được tạo.", "Order becomes debt/unpaid and debt record is created."),
        ("POST /orders/:id/debt thiếu customerName", "POST debt missing customerName", "API validation xử lý rõ ràng.", "API validation handles it clearly."),
        ("Thanh toán order đã paid lần nữa", "Pay an already paid order again", "API chặn double payment hoặc không cộng doanh thu lần hai.", "API blocks double payment or avoids double revenue."),
    ])

    add("03_api_backend/04_debt_stats_api/api_debt_stats_test_cases.xlsx", "API-FIN", "API - Công nợ và thống kê", "API - Debts and Statistics", "Kiểm tra API công nợ, thống kê ngày/tháng/năm", "API công nợ/thống kê", "Backend API > /debts, /stats", api_pre, [
        ("GET /debts trả danh sách công nợ", "GET /debts returns debt list", "Mỗi debt có customer, amount, status và orderId.", "Each debt has customer, amount, status, and orderId."),
        ("PATCH/POST thu nợ chuyển trạng thái paid", "Collect debt changes status to paid", "Debt được đánh dấu đã thu và có paidAt nếu API hỗ trợ.", "Debt is marked paid and has paidAt if supported."),
        ("Thu nợ id không tồn tại", "Collect non-existing debt id", "Trả 404/lỗi phù hợp, dữ liệu không đổi.", "Returns 404/proper error and data is unchanged."),
        ("GET /stats?period=day", "GET daily stats", "Trả revenue/orders/debts cho ngày hiện tại.", "Returns revenue/orders/debts for current day."),
        ("GET /stats?period=month", "GET monthly stats", "Trả tổng tháng đúng theo hóa đơn paid.", "Returns monthly totals based on paid invoices."),
        ("GET /stats?period=year", "GET yearly stats", "Trả tổng năm đúng theo dữ liệu.", "Returns yearly totals based on data."),
        ("GET /stats period không hợp lệ", "GET stats with invalid period", "API dùng default an toàn hoặc trả lỗi rõ.", "API uses safe default or returns clear error.", "period=weekx"),
        ("Stats bỏ qua order open", "Stats ignore open orders", "Order open không làm tăng revenue.", "Open orders do not increase revenue."),
        ("Stats tính công nợ riêng", "Stats report debts separately", "Debt amount không nhập nhằng với revenue đã thu.", "Debt amount is separate from collected revenue."),
        ("Stats sau reset dữ liệu", "Stats after data reset", "Trả số liệu seed ổn định/0 theo dữ liệu reset.", "Returns stable seed/zero metrics after reset."),
        ("Số tiền lớn trong stats", "Large amounts in stats", "Không overflow số hoặc format sai JSON.", "No numeric overflow or invalid JSON formatting."),
        ("GET /debts khi danh sách rỗng", "GET /debts when list is empty", "Trả mảng rỗng, không null bất ngờ.", "Returns empty array, not unexpected null."),
        ("Debt status chỉ dùng giá trị hợp lệ", "Debt status only uses valid values", "Status thuộc unpaid/paid/collected theo quy tắc.", "Status belongs to valid enum per rule."),
        ("Stats nhất quán sau nhiều request liên tiếp", "Stats remain consistent across repeated requests", "Nhiều request GET không làm đổi dữ liệu.", "Multiple GET requests do not mutate data."),
    ])

    add("03_api_backend/05_employee_attendance_api/api_employee_attendance_test_cases.xlsx", "API-HR", "API - Nhân viên và chấm công", "API - Employees and Attendance", "Kiểm tra API nhân viên, chấm công và payroll", "API nhân viên/chấm công", "Backend API > /employees, /attendance, /payroll", api_pre, [
        ("GET /employees trả danh sách nhân viên", "GET /employees returns employee list", "Mảng nhân viên có id, name, role, phone, salaryPerDay, active.", "Employee array has id, name, role, phone, salaryPerDay, active."),
        ("POST /employees thêm nhân viên", "POST /employees creates employee", "Nhân viên mới lưu với id không trùng.", "New employee is saved with unique id."),
        ("PATCH /employees/:id sửa nhân viên", "PATCH /employees/:id updates employee", "Nhân viên đúng id cập nhật, dữ liệu khác giữ nguyên.", "Target employee updates and other data remains unchanged."),
        ("DELETE/deactivate nhân viên", "Delete/deactivate employee", "Nhân viên inactive hoặc bị xóa theo thiết kế, không phá attendance cũ.", "Employee is inactive/deleted per design without breaking old attendance."),
        ("POST /attendance check-in", "POST /attendance check-in", "Record có employeeId, date, checkIn.", "Record has employeeId, date, checkIn."),
        ("PATCH/POST /attendance check-out", "Attendance check-out", "Record cập nhật checkOut và số giờ/ngày công.", "Record updates checkOut and worked time/days."),
        ("Check-in employee không tồn tại", "Check in non-existing employee", "API trả lỗi và không tạo record mồ côi.", "API returns error and creates no orphan record.", "employeeId=9999"),
        ("Check-in trùng ngày", "Duplicate check-in on same date", "API chặn hoặc merge record an toàn.", "API blocks or safely merges duplicate record."),
        ("Check-out không có check-in", "Check-out without check-in", "API trả lỗi hợp lệ.", "API returns valid error."),
        ("GET /attendance lọc theo ngày", "GET /attendance filtered by date", "Chỉ trả record đúng ngày.", "Only records for the selected date are returned."),
        ("GET /payroll mặc định tháng hiện tại", "GET /payroll current month default", "Trả bảng lương theo tháng hiện tại.", "Returns payroll for current month."),
        ("GET /payroll?month=YYYY-MM", "GET payroll by selected month", "Số công/lương khớp attendance trong tháng.", "Work days/salary match attendance in selected month."),
        ("Payroll bỏ qua nhân viên inactive nếu chính sách áp dụng", "Payroll excludes inactive employees if policy applies", "Nhân viên inactive không phát sinh lương mới.", "Inactive employees do not accrue new payroll."),
        ("salaryPerDay không phải số", "Non-numeric salaryPerDay", "API không lưu NaN gây hỏng payroll.", "API does not save NaN that breaks payroll.", "salaryPerDay=abc"),
        ("salaryPerDay âm", "Negative salaryPerDay", "API chặn hoặc chuẩn hóa lương âm.", "API blocks or normalizes negative salary.", "-500000"),
        ("Payroll nhiều request không mutate dữ liệu", "Repeated payroll requests do not mutate data", "GET payroll chỉ đọc, không tạo attendance mới.", "GET payroll is read-only and creates no attendance."),
    ])

    add("03_api_backend/06_admin_api/api_admin_test_cases.xlsx", "API-ADMIN", "API - Quản trị hệ thống", "API - System Admin", "Kiểm tra health, export và reset backend", "API admin", "Backend API > /admin, /reset", api_pre, [
        ("GET /admin/health trả status ok", "GET /admin/health returns status ok", "Response có status ok, storage, counts và generatedAt.", "Response includes ok status, storage, counts, and generatedAt."),
        ("GET /admin/export trả dữ liệu JSON", "GET /admin/export returns JSON data", "Export có cấu trúc users/tables/menu/orders/debts/employees.", "Export has users/tables/menu/orders/debts/employees."),
        ("Export không làm thay đổi dữ liệu", "Export does not mutate data", "GET export nhiều lần trả dữ liệu ổn định nếu không có thao tác ghi.", "Repeated export returns stable data if no writes occur."),
        ("POST /admin/reset khôi phục seed", "POST /admin/reset restores seed data", "Tables/menu/users/nextIds quay về seed.", "Tables/menu/users/nextIds return to seed."),
        ("POST /reset root khôi phục seed", "POST /reset root restores seed data", "Endpoint root reset hoạt động tương đương hoặc theo thiết kế.", "Root reset behaves equivalently or per design."),
        ("Reset sau khi có order/debt", "Reset after orders/debts exist", "Orders/debts bị xóa và counts về seed.", "Orders/debts are removed and counts return to seed."),
        ("Health sau khi reset", "Health after reset", "Counts đúng 6 bàn, 5 món, 3 nhân viên theo seed.", "Counts match six tables, five menu items, and three employees."),
        ("Admin endpoint trả JSON hợp lệ khi lỗi đọc file", "Admin endpoint returns valid JSON on file read error", "Server xử lý lỗi file an toàn, không crash process.", "Server handles file errors safely without crashing."),
        ("CORS cho web app localhost", "CORS allows localhost web app", "Web gọi API không bị chặn CORS trong dev.", "Web can call API without CORS block in dev."),
        ("Server lắng nghe cổng 3000", "Server listens on port 3000", "Port 3000 phản hồi API root/admin đúng.", "Port 3000 responds to root/admin API."),
    ])

    add("04_cross_platform/01_end_to_end/e2e_business_flow_test_cases.xlsx", "E2E-FLOW", "End-to-end - Luồng nghiệp vụ", "End-to-end - Business Flows", "Kiểm tra luồng thực tế liên kết mobile, web và backend", "Luồng nghiệp vụ", "Mobile + Web + Backend", e2e_pre, [
        ("Admin đăng nhập, thêm bàn, mở bàn trên mobile", "Admin logs in, adds a table, and opens it on mobile", "Bàn mới xuất hiện trên mobile và backend.", "New table appears on mobile and backend."),
        ("Mobile thêm món tạo hóa đơn, web thấy hóa đơn", "Mobile adds item and web sees invoice", "Hóa đơn tạo trên mobile xuất hiện ở web sau refresh.", "Invoice created on mobile appears on web after refresh."),
        ("Web thêm món tiếp vào hóa đơn mobile đã tạo", "Web adds another item to mobile-created invoice", "Cùng hóa đơn cập nhật item/tổng tiền trên cả hai nền tảng.", "Same invoice updates item/total on both platforms."),
        ("Admin thanh toán trên web, mobile cập nhật trạng thái", "Admin pays on web and mobile updates status", "Hóa đơn paid và bàn trống sau refresh mobile.", "Invoice is paid and table is available after mobile refresh."),
        ("Staff tạo hóa đơn, admin duyệt thanh toán", "Staff creates invoice and admin completes payment", "Nhân viên ghi nhận đúng, admin thanh toán thành công.", "Staff is recorded correctly and admin payment succeeds."),
        ("Ghi nợ trên mobile, thu nợ trên web", "Create debt on mobile and collect on web", "Debt status đồng bộ qua hai nền tảng.", "Debt status syncs across platforms."),
        ("Admin thêm món trên web, mobile gọi món thấy món mới", "Admin adds menu item on web and mobile order entry sees it", "Món mới available xuất hiện ở màn gọi món mobile.", "New available item appears in mobile order entry."),
        ("Admin tắt bán món, web/mobile đều ẩn ở gọi món", "Admin disables item and both web/mobile hide it in order entry", "Món unavailable không được thêm mới trên cả hai.", "Unavailable item cannot be newly added on both platforms."),
        ("Chấm công trên mobile, payroll web cập nhật", "Attendance on mobile updates payroll on web", "Bảng lương phản ánh ngày công mới.", "Payroll reflects the new attendance day."),
        ("Admin tạo thông báo trên web, staff đọc trên mobile", "Admin creates notice on web and staff reads on mobile", "Thông báo mới xuất hiện với nội dung đúng.", "New notice appears with correct content."),
        ("Reset demo từ web admin rồi mobile refresh", "Reset demo from web admin then refresh mobile", "Mobile trở về dữ liệu seed không còn order/debt cũ.", "Mobile returns to seed data without old orders/debts."),
        ("Mất mạng backend khi đang dùng cả web và mobile", "Backend outage while both web and mobile are active", "Cả hai nền tảng báo lỗi thân thiện và không ghi dữ liệu nửa chừng.", "Both platforms show friendly errors and avoid partial writes."),
        ("Hai người cùng mở một bàn", "Two users open the same table", "Hệ thống không tạo hai hóa đơn open cho một bàn.", "System does not create two open invoices for one table."),
        ("Hóa đơn paid không thể chỉnh món thêm", "Paid invoice cannot receive more items", "API/UI chặn sửa đơn đã chốt.", "API/UI blocks edits to closed invoice."),
        ("Dữ liệu thống kê sau luồng bán hàng đầy đủ", "Statistics after a complete sale flow", "Doanh thu/ngày công/công nợ khớp các thao tác đã làm.", "Revenue/attendance/debts match performed actions."),
        ("Khởi động lại backend không mất dữ liệu JSON", "Restart backend without losing JSON data", "Dữ liệu đã lưu vẫn tồn tại sau restart.", "Saved data persists after restart."),
    ])

    add("04_cross_platform/02_permissions_security/permissions_security_test_cases.xlsx", "SEC-ROLE", "Bảo mật - Quyền và dữ liệu", "Security - Roles and Data", "Kiểm tra quyền admin/staff, dữ liệu nhạy cảm và thao tác nguy hiểm", "Quyền và bảo mật", "Mobile/Web/API", e2e_pre, [
        ("Admin có quyền quản lý nhân viên", "Admin can manage employees", "Admin thấy và thao tác CRUD nhân viên theo thiết kế.", "Admin can view and use employee CRUD per design."),
        ("Staff không có quyền quản lý nhân viên", "Staff cannot manage employees", "Staff bị ẩn menu hoặc bị API từ chối.", "Staff menu is hidden or API rejects access."),
        ("Admin có quyền tính lương toàn bộ", "Admin can view full payroll", "Bảng lương toàn bộ chỉ mở cho admin.", "Full payroll is available only to admin."),
        ("Staff không xem lương toàn bộ", "Staff cannot view all payroll", "Staff không xem dữ liệu lương người khác.", "Staff cannot view other employees payroll."),
        ("Admin/staff được quyền thu tiền hóa đơn theo yêu cầu", "Admin/staff can collect invoice per requirement", "Nút Đã trả hiển thị cho admin hoặc nhân viên được phép.", "Paid button is shown for admin or allowed staff."),
        ("Role lạ không thanh toán được hóa đơn", "Unknown role cannot collect invoice", "API trả lỗi unauthorized/forbidden.", "API returns unauthorized/forbidden.", "role=guest"),
        ("Password không xuất hiện trong response login", "Password does not appear in login response", "Response user không có trường password.", "Login response user has no password field."),
        ("Password không xuất hiện trong export quản trị", "Password does not appear in admin export", "Export loại bỏ/mask password nếu áp dụng bảo mật.", "Export removes/masks password when security policy applies."),
        ("Input text không gây lỗi render", "Text input does not break rendering", "Tên bàn/món/nhân viên chứa ký tự đặc biệt vẫn hiển thị an toàn.", "Special characters in names render safely.", "<script>alert(1)</script>"),
        ("Giá trị số âm bị chặn ở tiền và lương", "Negative numeric values are blocked for money and salary", "Không lưu giá/lương/giảm giá âm bất hợp lý.", "Invalid negative price/salary/discount is not saved."),
        ("ID không tồn tại không làm crash API", "Non-existing ids do not crash API", "API trả 404/lỗi rõ, server vẫn sống.", "API returns clear 404/error and server keeps running."),
        ("Request JSON quá thiếu field", "JSON request with missing fields", "Server validate/normalize an toàn.", "Server validates/normalizes safely."),
        ("CORS chỉ dùng cho môi trường dev rõ ràng", "CORS behavior is clear for dev environment", "Web localhost gọi được API; production cần chính sách tương ứng.", "Localhost web can call API; production needs proper policy."),
        ("Không lưu dữ liệu nửa chừng khi API lỗi", "No partial save on API error", "Thao tác lỗi không tạo bản ghi thiếu liên kết.", "Failed operation does not create orphan/partial records."),
        ("Reset demo là thao tác nguy hiểm", "Demo reset is a dangerous action", "UI cần xác nhận; test đảm bảo không bấm nhầm gây mất dữ liệu.", "UI requires confirmation to avoid accidental data loss."),
        ("Staff không truy cập trực tiếp URL web admin", "Staff cannot access web admin by direct URL/state", "Truy cập trực tiếp vẫn bị chặn hoặc redirect.", "Direct access is blocked or redirected."),
        ("Dữ liệu tiền không bị NaN/Infinity", "Money values never become NaN/Infinity", "API/UI không hiển thị hoặc lưu NaN/Infinity.", "API/UI never displays or saves NaN/Infinity."),
        ("Dữ liệu sau reset dùng password duyencocong2", "Seed data after reset uses password duyencocong2", "Admin/staff đăng nhập được bằng password mới sau reset.", "Admin/staff can log in with the new password after reset.", "admin/staff / duyencocong2"),
    ])

    add("04_cross_platform/03_usability_visual/usability_visual_test_cases.xlsx", "UX-VISUAL", "UI/UX - Hình ảnh và khả dụng", "UI/UX - Visual and Usability", "Kiểm tra phong cách đen trắng xanh nước biển, hình ảnh và thao tác người dùng", "UI/UX", "Mobile/Web UI", e2e_pre, [
        ("Tên app hiển thị là Quản Lý Nhà Hàng, không phải frontend", "App name is Restaurant Management, not frontend", "Android launcher/web title/app bar dùng đúng tên.", "Android launcher/web title/app bar use correct name."),
        ("Logo tham chiếu hiển thị đúng tỉ lệ", "Reference logo displays with correct ratio", "Logo không méo, không bị cắt xấu.", "Logo is not distorted or badly cropped."),
        ("Palette đen trắng xanh nước biển nhất quán", "Black-white-ocean-blue palette is consistent", "Nút chính, app bar, chip trạng thái dùng màu hợp lý.", "Primary buttons, app bars, and chips use coherent colors."),
        ("Màu trạng thái bàn dễ phân biệt", "Table status colors are distinguishable", "Trống/đang dùng/đặt trước có màu và nhãn rõ.", "Available/occupied/reserved have clear colors and labels."),
        ("Ảnh món ăn đẹp và không giống placeholder lỗi", "Food images are attractive and not broken placeholders", "Ảnh tải từ URL mạng hoặc fallback đẹp.", "Network images or fallbacks look good."),
        ("Text trên nút không bị cắt ở mobile", "Button text is not clipped on mobile", "Nút Thêm bàn, Gọi món, Đã trả đọc được ở màn nhỏ.", "Add table, order, and paid buttons remain readable on small screens."),
        ("Không có cảnh báo overflow trên app", "No overflow warning in app", "Không còn sọc vàng/đen ở card bàn hoặc màn khác.", "No yellow/black overflow stripes on table cards or other screens."),
        ("Form dialog vừa màn hình điện thoại", "Dialogs fit phone screen", "Dialog thêm/sửa bàn/món/nhân viên có thể cuộn nếu cần.", "Add/edit dialogs are scrollable when needed."),
        ("Loading state hiển thị trong lúc gọi API", "Loading state appears during API requests", "Spinner/skeleton xuất hiện và biến mất đúng lúc.", "Spinner/skeleton appears and disappears correctly."),
        ("Empty state rõ ràng khi chưa có dữ liệu", "Empty state is clear when no data exists", "Hóa đơn/công nợ/trống dữ liệu có thông báo dễ hiểu.", "Invoices/debts/empty data show clear messages."),
        ("Snackbar lỗi không che thao tác quan trọng quá lâu", "Error snackbar does not block important actions too long", "Lỗi đọc được nhưng không làm app kẹt.", "Error is readable without trapping the user."),
        ("Các chức năng chính thao tác được bằng một tay trên mobile", "Primary workflows are reachable one-handed on mobile", "Nút chính nằm trong vùng dễ bấm và kích thước đủ lớn.", "Primary actions are reachable and large enough."),
    ])

    return specs


def create_master_summary() -> Path:
    path = ROOT / "00_summary" / "00_master_test_report.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Cover"
    ws.merge_cells("A1:H2")
    ws["A1"] = "MASTER TEST REPORT"
    ws["A1"].font = Font(size=22, bold=True, color="FFFFFF")
    ws["A1"].fill = header_fill
    ws["A1"].alignment = center
    info = [
        ("Project Name", PROJECT_NAME),
        ("Project Code", PROJECT_CODE),
        ("Version", VERSION),
        ("Issue date", ISSUE_DATE),
        ("Total workbooks", len(workbook_summaries)),
        ("Total test cases", sum(int(x["count"]) for x in workbook_summaries)),
        ("Total Pass", sum(int(x["pass"]) for x in workbook_summaries)),
        ("Total Fail", sum(int(x["fail"]) for x in workbook_summaries)),
        ("Total Pending", sum(int(x["pending"]) for x in workbook_summaries)),
    ]
    for row, (label, value) in enumerate(info, start=4):
        ws.cell(row, 1, label).font = bold_font
        ws.cell(row, 1).fill = meta_fill
        ws.cell(row, 2, value)
        ws.cell(row, 1).border = border
        ws.cell(row, 2).border = border
    for idx, width in enumerate([22, 48, 18, 18, 18, 18, 18, 18], 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    detail = wb.create_sheet("Workbook Summary")
    headers = ["No.", "Area / Khu vực", "Module Code", "Module (VI)", "Module (EN)", "File path", "Cases", "Pass", "Fail", "Pending"]
    for col, h in enumerate(headers, 1):
        cell = detail.cell(1, col, h)
        cell.fill = header_fill
        cell.font = white_font
        cell.border = border
        cell.alignment = center
    for row, item in enumerate(workbook_summaries, start=2):
        parts = str(item["path"]).split("/")
        area = parts[2] if len(parts) > 2 else ""
        values = [
            row - 1,
            area,
            item["module_code"],
            item["module_vi"],
            item["module_en"],
            item["path"],
            item["count"],
            item["pass"],
            item["fail"],
            item["pending"],
        ]
        for col, value in enumerate(values, 1):
            cell = detail.cell(row, col, value)
            cell.border = border
            cell.alignment = wrap if col in (4, 5, 6) else center
    widths = [8, 24, 18, 34, 34, 78, 12, 12, 12, 12]
    for idx, width in enumerate(widths, 1):
        detail.column_dimensions[get_column_letter(idx)].width = width
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = f"A1:J{len(workbook_summaries)+1}"

    report = wb.create_sheet("Test Report")
    report.merge_cells("A1:G1")
    report["A1"] = "MASTER TEST REPORT"
    report["A1"].fill = header_fill
    report["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    report["A1"].alignment = center
    rows = [
        ("Number of workbook files", len(workbook_summaries)),
        ("Number of test cases", sum(int(x["count"]) for x in workbook_summaries)),
        ("Pass", sum(int(x["pass"]) for x in workbook_summaries)),
        ("Fail", sum(int(x["fail"]) for x in workbook_summaries)),
        ("Pending", sum(int(x["pending"]) for x in workbook_summaries)),
        ("Duplicate IDs", 0),
    ]
    for row, (label, value) in enumerate(rows, start=3):
        report.cell(row, 1, label).font = bold_font
        report.cell(row, 1).fill = meta_fill
        report.cell(row, 2, value)
        report.cell(row, 1).border = border
        report.cell(row, 2).border = border
    for idx, width in enumerate([28, 18, 28, 28, 28, 28, 28], 1):
        report.column_dimensions[get_column_letter(idx)].width = width
    wb.save(path)
    return path


def main() -> None:
    smoke = executed_smoke_cases()
    save_workbook("00_summary/01_executed_smoke_tests.xlsx", "SMK", "Smoke đã chạy", "Executed Smoke Tests", "Các kiểm tra thật đã chạy và ghi kết quả Pass", smoke)
    for rel, code, vi, en, req, cases in pending_specs():
        save_workbook(rel, code, vi, en, req, cases)

    ids = [item[0] for item in all_registry]
    duplicates = [id_ for id_, count in Counter(ids).items() if count > 1]
    if duplicates:
        raise SystemExit(f"Duplicate test case IDs found: {duplicates[:20]}")

    validation = c(
        "SMK-010",
        "Smoke verification",
        "Document/test-cases",
        "Kiểm tra toàn bộ ID test case không trùng nhau",
        "Verify all generated test case IDs are unique",
        "Không có ID trùng trong toàn bộ bộ tài liệu.",
        "No duplicate IDs exist across the documentation set.",
        data=f"{len(ids)} IDs checked before validation case",
        preconditions="Bộ test case đã được sinh bằng script openpyxl.",
        result=PASS,
        actual="0 duplicate IDs found before adding validation row; final total includes SMK-010.",
        test_date=ISSUE_DATE,
    )

    for item in list(workbook_summaries):
        if item["module_code"] == "SMK":
            workbook_summaries.remove(item)
    all_registry[:] = [entry for entry in all_registry if not entry[0].startswith("SMK-")]
    smoke.append(validation)
    save_workbook("00_summary/01_executed_smoke_tests.xlsx", "SMK", "Smoke đã chạy", "Executed Smoke Tests", "Các kiểm tra thật đã chạy và ghi kết quả Pass", smoke)

    ids = [item[0] for item in all_registry]
    duplicates = [id_ for id_, count in Counter(ids).items() if count > 1]
    if duplicates:
        raise SystemExit(f"Duplicate test case IDs found after validation: {duplicates[:20]}")

    master_path = create_master_summary()

    created = sorted(ROOT.rglob("*.xlsx"))
    for xlsx in created:
        wb = load_workbook(xlsx, read_only=True, data_only=False)
        required = {"Cover", "Test Report"}
        if not required.issubset(set(wb.sheetnames)):
            raise SystemExit(f"Missing required sheets in {xlsx}")

    print(f"created_files={len(created)}")
    print(f"test_case_files={len(workbook_summaries)}")
    print(f"master={master_path.as_posix()}")
    print(f"total_cases={sum(int(x['count']) for x in workbook_summaries)}")
    print(f"pass={sum(int(x['pass']) for x in workbook_summaries)}")
    print(f"fail={sum(int(x['fail']) for x in workbook_summaries)}")
    print(f"pending={sum(int(x['pending']) for x in workbook_summaries)}")
    print(f"duplicate_ids={len([id_ for id_, count in Counter(ids).items() if count > 1])}")


if __name__ == "__main__":
    main()
