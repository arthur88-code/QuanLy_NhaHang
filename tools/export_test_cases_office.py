from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Document" / "test-cases" / "all-test-cases-expandtesting-style.md"
DOCX_OUT = ROOT / "Document" / "test-cases" / "all-test-cases-expandtesting-style.docx"
XLSX_OUT = ROOT / "Document" / "test-cases" / "all-test-cases-expandtesting-style.xlsx"


def parse_cases(text: str) -> list[dict[str, str]]:
    sections = re.split(r"\n### Test Case ", text)
    cases: list[dict[str, str]] = []
    current_group = ""
    lines = text.splitlines()
    group_by_number: dict[int, str] = {}
    group = ""
    for line in lines:
      if line.startswith("## ") and line.endswith(" Automation Test Cases"):
          group = line[3:].replace(" Automation Test Cases", "")
      if line.startswith("### Test Case "):
          match = re.match(r"### Test Case (\d+):", line)
          if match:
              group_by_number[int(match.group(1))] = group

    for raw in sections[1:]:
        header, _, body = raw.partition("\n")
        number_text, _, title = header.partition(": ")
        number = int(number_text.strip())
        current_group = group_by_number.get(number, current_group)
        test_id = re.search(r"\*\*Test ID:\*\* (.+)", body)
        location = re.search(r"\*\*Test Location:\*\* (.+)", body)
        preconditions = re.search(r"\*\*Preconditions:\*\* (.+)", body)
        expected = re.search(r"\*\*Expected Result:\*\*\n((?:- .+\n?)+)", body)
        status = re.search(r"\*\*Status:\*\* (.+)", body)
        steps_match = re.search(r"\*\*Steps:\*\*\n(.+?)\n\n\*\*Expected Result:\*\*", body, re.S)
        expected_text = ""
        if expected:
            expected_text = "\n".join(line[2:] for line in expected.group(1).splitlines() if line.startswith("- "))
        cases.append(
            {
                "No": str(number),
                "Group": current_group,
                "Title": title.strip(),
                "Test ID": test_id.group(1).strip() if test_id else "",
                "Location": location.group(1).strip() if location else "",
                "Preconditions": preconditions.group(1).strip() if preconditions else "",
                "Steps": steps_match.group(1).strip() if steps_match else "",
                "Expected Result": expected_text,
                "Status": status.group(1).strip() if status else "",
            }
        )
    return cases


def export_docx(cases: list[dict[str, str]]) -> None:
    doc = Document()
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    doc.add_heading("Restaurant Management System - Test Cases", level=0)
    doc.add_paragraph(f"Source: {SOURCE.name}")
    doc.add_paragraph(f"Total test cases: {len(cases)}")
    last_group = ""
    for case in cases:
        if case["Group"] != last_group:
            doc.add_heading(f"{case['Group']} Automation Test Cases", level=1)
            last_group = case["Group"]
        doc.add_heading(f"Test Case {case['No']}: {case['Title']}", level=2)
        for label in ["Test ID", "Location", "Preconditions", "Steps", "Expected Result", "Status"]:
            paragraph = doc.add_paragraph()
            paragraph.add_run(f"{label}: ").bold = True
            paragraph.add_run(case[label])
    doc.save(DOCX_OUT)


def export_xlsx(cases: list[dict[str, str]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Cases"
    headers = ["No", "Group", "Test ID", "Title", "Location", "Preconditions", "Steps", "Expected Result", "Status"]
    fill = PatternFill("solid", fgColor="0B5B86")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)
        cell.fill = fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row, case in enumerate(cases, 2):
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row, col, case[header])
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [8, 30, 16, 42, 42, 48, 56, 56, 22]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{len(cases) + 1}"
    wb.save(XLSX_OUT)


def main() -> None:
    cases = parse_cases(SOURCE.read_text(encoding="utf-8"))
    export_docx(cases)
    export_xlsx(cases)
    print(f"cases={len(cases)}")
    print(f"docx={DOCX_OUT}")
    print(f"xlsx={XLSX_OUT}")


if __name__ == "__main__":
    main()
